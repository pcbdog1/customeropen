from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import Lead
from .utils import dedupe_key

LEADS_SHEET = "Leads"
DUPLICATE_SHEET = "Duplicate_Skipped"

HEADERS = [
    "序号",
    "添加日期",
    "国家",
    "城市",
    "公司名称",
    "公司官网",
    "官网域名",
    "主营业务",
    "所属行业",
    "为什么判断它可能需要 PCB/PCBA",
    "需求信号类型",
    "需求信号来源",
    "公司人数",
    "成立时间",
    "公开邮箱",
    "Email Type",
    "Email Source URL",
    "Contact Method",
    "Contact Form URL",
    "联系人姓名",
    "联系人职位",
    "联系电话",
    "LinkedIn 公司主页",
    "数据来源链接",
    "客户匹配评分",
    "推荐开发角度",
    "建议开发邮件主题",
    "备注",
    "Compliance Note",
    "跟进状态",
    "最近跟进日期",
    "邮件发送状态",
    "去重键",
]


@dataclass
class AppendResult:
    added: int
    skipped_duplicates: int


class ExcelLeadStore:
    def __init__(self, path: str | Path):
        self.path = Path(path)

    def _workbook(self):
        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = LEADS_SHEET
            ws.append(HEADERS)
            return wb
        wb = load_workbook(self.path)
        if LEADS_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(LEADS_SHEET)
            ws.append(HEADERS)
        else:
            self._ensure_headers(wb[LEADS_SHEET])
        return wb

    def _ensure_headers(self, ws) -> None:
        existing = [cell.value for cell in ws[1]]
        for header in HEADERS:
            if header not in existing:
                ws.cell(row=1, column=len(existing) + 1, value=header)
                existing.append(header)

    def load_existing_keys(self) -> set[str]:
        wb = self._workbook()
        ws = wb[LEADS_SHEET]
        headers = [cell.value for cell in ws[1]]
        try:
            name_idx = headers.index("公司名称") + 1
            website_idx = headers.index("公司官网") + 1
        except ValueError:
            return set()

        keys: set[str] = set()
        for row in range(2, ws.max_row + 1):
            company = ws.cell(row=row, column=name_idx).value or ""
            website = ws.cell(row=row, column=website_idx).value or ""
            if company or website:
                keys.add(dedupe_key(str(company), str(website)))
        return keys

    def append_leads(self, leads: list[Lead]) -> AppendResult:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = self._workbook()
        ws = wb[LEADS_SHEET]
        existing_keys = self.load_existing_keys()
        duplicate_rows: list[dict[str, object]] = []
        added = 0

        for lead in leads:
            key = dedupe_key(lead.company_name, lead.website)
            if key in existing_keys:
                duplicate_rows.append({"公司名称": lead.company_name, "公司官网": lead.website, "去重键": key})
                continue
            existing_keys.add(key)
            added += 1
            row_data = lead.to_excel_row(ws.max_row)
            ws.append([row_data.get(header, "") for header in HEADERS])

        if DUPLICATE_SHEET in wb.sheetnames:
            del wb[DUPLICATE_SHEET]
        dup_ws = wb.create_sheet(DUPLICATE_SHEET)
        dup_ws.append(["公司名称", "公司官网", "去重键"])
        for row in duplicate_rows:
            dup_ws.append([row["公司名称"], row["公司官网"], row["去重键"]])

        if "Search_Log" in wb.sheetnames:
            del wb["Search_Log"]
        log_ws = wb.create_sheet("Search_Log")
        log_ws.append(["字段", "值"])
        log_ws.append(["最后新增数量", added])
        log_ws.append(["跳过重复数量", len(duplicate_rows)])
        log_ws.append(["总客户数量", ws.max_row - 1])

        wb.save(self.path)
        return AppendResult(added=added, skipped_duplicates=len(duplicate_rows))
