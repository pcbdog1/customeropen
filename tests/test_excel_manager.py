from __future__ import annotations

import pytest
from openpyxl import Workbook

from email_sender_common import OPT_OUT_SENTENCE
from excel_manager import OUTREACH_SHEET, SENT_LOG_SHEET, ExcelManager
from pcb_leads.excel_store import HEADERS as LEAD_HEADERS


@pytest.fixture
def manager(tmp_path):
    path = tmp_path / "leads.xlsx"
    workbook = Workbook()
    workbook.active.title = "Leads"
    workbook.save(path)
    return ExcelManager(path)


def draft(company: str, email: str, domain: str, country: str) -> dict[str, str]:
    return {
        "Company Name": company,
        "Contact Email": email,
        "Company Domain": domain,
        "Country": country,
    }


def test_contact_titles_on_different_domains_do_not_collide(manager):
    first = draft("Contact", "sales@alpha.example", "alpha.example", "Germany")
    second = draft("Contact", "sales@beta.example", "beta.example", "Germany")

    assert manager.append_outreach_drafts([first, second]) == 2


def test_about_titles_on_different_domains_do_not_collide(manager):
    first = draft("About", "sales@alpha.example", "alpha.example", "Germany")
    second = draft("About", "sales@beta.example", "beta.example", "Germany")

    assert manager.append_outreach_drafts([first, second]) == 2


def test_same_email_is_duplicate(manager):
    first = draft("Alpha", "sales@alpha.example", "alpha.example", "Germany")
    second = draft("Alpha GmbH", "sales@alpha.example", "alpha-gmbh.example", "Germany")

    assert manager.append_outreach_drafts([first, second]) == 1


def test_same_complete_hostname_is_duplicate(manager):
    first = draft("Alpha", "sales@alpha.example", "shop.alpha.example", "Germany")
    second = draft("Beta", "sales@beta.example", "shop.alpha.example", "Germany")

    assert manager.append_outreach_drafts([first, second]) == 1


def test_normalized_company_and_country_is_fallback_duplicate(manager):
    first = draft("Alpha Controls GmbH", "sales@alpha.de", "alpha.de", "Germany")
    second = draft(" Alpha Controls GmbH ", "office@alpha-controls.de", "alpha-controls.de", " Germany ")

    assert manager.append_outreach_drafts([first, second]) == 1


def test_sent_log_email_company_and_domain_are_permanent_exclusions(manager):
    manager.sheet(SENT_LOG_SHEET).append(
        [
            "2026-08-07T10:00:00",
            "Sent Alpha",
            "sales@sent-alpha.example",
            "subject",
            "SENT",
            "",
            "Germany",
            "Automation",
            "https://sent-alpha.example",
            "sent-alpha.example",
        ]
    )

    assert manager.append_outreach_drafts(
        [
            draft("New Company", "sales@sent-alpha.example", "new.example", "Germany"),
            draft("New Company", "sales@new.example", "sent-alpha.example", "Germany"),
            draft("Sent Alpha", "sales@another.example", "another.example", "Germany"),
        ]
    ) == 0


def test_sent_domains_derives_legacy_blank_domain_from_attributable_source_and_email(manager):
    manager.sheet(SENT_LOG_SHEET).append(
        [
            "2026-08-07T10:00:00",
            "Legacy Alpha",
            "sales@legacy-alpha.example",
            "subject",
            "SENT",
            "",
            "Germany",
            "Automation",
            "https://www.legacy-alpha.example/contact",
            "",
        ]
    )

    assert manager.sent_domains() == {"legacy-alpha.example"}


def test_final_send_safety_blocks_wikipedia_category_and_unrelated_email_source(manager):
    ws = manager.sheet("Outreach_Email_Drafts")
    body = " ".join(["word"] * 125) + "\n\nIf you are not the right person or prefer not to receive further emails from us, please let me know and I will remove your contact from our list."
    ws.append([
        "Category: Medical technology companies of Sweden - Wikipedia", "lisa@wikipedia.org",
        "Public Work Email", "Subject", body, "hardware signals: sensor, controller.", "", "", "", "", "",
        "https://en.wikipedia.org/wiki/Category:Medical_technology_companies_of_Sweden", "note", "", 4,
        "wikipedia.org", "Sweden", "Medical", 125, "Medical PCBA",
    ])
    manager.evaluate_auto_send_all()
    headers = manager.headers(ws)
    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert ws.cell(2, headers.index("Final Send Safety Check") + 1).value == "BLOCKED"
    assert "Blocked encyclopedia/wiki source" in ws.cell(2, headers.index("Final Send Block Reason") + 1).value


def test_final_send_safety_blocks_sentence_style_page_title_as_company_name(manager):
    ws = manager.sheet("Outreach_Email_Drafts")
    body = " ".join(["word"] * 125) + "\n\nIf you are not the right person or prefer not to receive further emails from us, please let me know and I will remove your contact from our list."
    ws.append([
        "Anaesthetic devices is Breas Medical AB", "info@vinnova.example", "Generic Company Email", "Subject", body,
        "hardware signals: sensor, controller.", "", "", "", "", "", "https://vinnova.example/contact", "note", "", 4,
        "vinnova.example", "Sweden", "Medical", 125, "Medical PCBA",
    ])
    manager.evaluate_auto_send_all()
    headers = manager.headers(ws)
    assert ws.cell(2, headers.index("Final Send Safety Check") + 1).value == "BLOCKED"
    assert "page title or sentence" in ws.cell(2, headers.index("Final Send Block Reason") + 1).value


def test_final_send_safety_accepts_attributable_company_contact_source(manager):
    ws = manager.sheet("Outreach_Email_Drafts")
    body = " ".join(["word"] * 125) + "\n\nIf you are not the right person or prefer not to receive further emails from us, please let me know and I will remove your contact from our list."
    ws.append([
        "Acme Controls GmbH", "sales@acme.example", "Generic Company Email", "Subject", body,
        "hardware signals: sensor, controller.", "", "", "", "", "", "https://acme.example/contact", "note", "", 4,
        "acme.example", "Germany", "Automation", 125, "Industrial Control PCB",
    ])
    manager.evaluate_auto_send_all()
    headers = manager.headers(ws)
    assert ws.cell(2, headers.index("Final Send Safety Check") + 1).value == "PASS"


def test_final_send_safety_allows_directory_discovery_when_company_official_source_is_present(manager):
    ws = manager.sheet("Outreach_Email_Drafts")
    body = " ".join(["word"] * 125) + "\n\nIf you are not the right person or prefer not to receive further emails from us, please let me know and I will remove your contact from our list."
    ws.append([
        "Acme Controls GmbH", "sales@acme.example", "Generic Company Email", "Subject", body,
        "hardware signals: sensor, controller.", "", "", "", "", "",
        "https://exhibitor-directory.example/acme; https://acme.example/impressum", "note", "", 4,
        "acme.example", "Germany", "Automation", 125, "Industrial Control PCB",
    ])
    manager.evaluate_auto_send_all()
    headers = manager.headers(ws)
    assert ws.cell(2, headers.index("Final Send Safety Check") + 1).value == "PASS"


@pytest.mark.parametrize(
    ("email", "source"),
    [
        ("sales@yahoo.com", "https://yahoo.com/contact"),
        ("sales@legacy-alpha.example", "not a valid source url"),
        ("sales@legacy-alpha.example", "https://unrelated.example/contact"),
        ("not-an-email", "https://legacy-alpha.example/contact"),
    ],
)
def test_sent_domains_rejects_free_malformed_or_unrelated_legacy_derivations(manager, email, source):
    manager.sheet(SENT_LOG_SHEET).append(
        [
            "2026-08-07T10:00:00",
            "Legacy Alpha",
            email,
            "subject",
            "SENT",
            "",
            "Germany",
            "Automation",
            source,
            "",
        ]
    )

    assert manager.sent_domains() == set()


@pytest.mark.parametrize(
    "email",
    [
        "sales@@acme.co.uk",
        ".sales@acme.co.uk",
        "sales.@acme.co.uk",
        "sa..les@acme.co.uk",
        "sales@-acme.co.uk",
        "sales@acme..co.uk",
        "sales@acme.co.uk:25",
        "sales@co.uk",
        "\u017fales@acme.co.uk",
        "sales@acme\u2024co.uk",
        "sales@acme\ufe52co.uk",
        "sales@b\u00fccher.de",
        "sales@.acme.co.uk",
        "sales@acme.co.uk.",
        "sales@www.acme.co.uk",
        "sales@acme-.co.uk",
        "sales@ac_me.co.uk",
    ],
)
def test_sent_domains_uses_strict_full_mailbox_validation_for_legacy_rows(manager, email):
    manager.sheet(SENT_LOG_SHEET).append(
        [
            "2026-08-07T10:00:00",
            "Legacy Acme",
            email,
            "subject",
            "SENT",
            "",
            "UK",
            "Automation",
            "https://acme.co.uk/contact",
            "",
        ]
    )

    assert manager.sent_domains() == set()


def append_lead_row(manager: ExcelManager, **values: str) -> None:
    manager.sheet("Leads").append([values.get(header, "") for header in LEAD_HEADERS])


def test_evidence_backfill_prefers_matching_complete_email_domain(manager):
    append_lead_row(
        manager,
        **{
            "公司名称": "Shared Controls",
            "公司官网": "https://alpha.example",
            "公开邮箱": "sales@alpha.example",
            "国家": "Germany",
            "为什么判断它可能需要 PCB/PCBA": "Alpha evidence. hardware signals: sensor, controller.",
        },
    )
    append_lead_row(
        manager,
        **{
            "公司名称": "Shared Controls",
            "公司官网": "https://beta.example",
            "公开邮箱": "sales@beta.example",
            "国家": "Germany",
            "为什么判断它可能需要 PCB/PCBA": "Beta evidence. hardware signals: battery, power electronics.",
        },
    )
    ws = manager.sheet(OUTREACH_SHEET)
    ws.append(["Shared Controls", "sales@alpha.example"])

    manager.backfill_outreach_metadata()

    headers = manager.headers(ws)
    assert ws.cell(2, headers.index("Company Domain") + 1).value == "alpha.example"
    assert "Alpha evidence" in ws.cell(2, headers.index("Why This Angle Fits") + 1).value
    assert "Beta evidence" not in ws.cell(2, headers.index("Why This Angle Fits") + 1).value


def test_evidence_backfill_uses_company_country_only_when_unambiguous(manager):
    append_lead_row(
        manager,
        **{
            "公司名称": "Unique Instruments GmbH",
            "国家": "Germany",
            "为什么判断它可能需要 PCB/PCBA": "Unique evidence. hardware signals: instrument, measurement.",
        },
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)
    ws.append(["Unique Instruments GmbH", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "Germany"])

    manager.backfill_outreach_metadata()

    assert "Unique evidence" in ws.cell(2, headers.index("Why This Angle Fits") + 1).value


def test_evidence_backfill_does_not_use_ambiguous_company_country(manager):
    for domain, evidence in (("alpha.example", "Alpha evidence"), ("beta.example", "Beta evidence")):
        append_lead_row(
            manager,
            **{
                "公司名称": "Shared Controls",
                "公司官网": f"https://{domain}",
                "国家": "Germany",
                "为什么判断它可能需要 PCB/PCBA": f"{evidence}. hardware signals: sensor, controller.",
            },
        )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)
    ws.append(["Shared Controls", "", "", "", "", "Original angle", "", "", "", "", "", "", "", "", "", "", "Germany"])

    manager.backfill_outreach_metadata()

    assert ws.cell(2, headers.index("Why This Angle Fits") + 1).value == "Original angle"


def sendable_draft(
    company: str = "Acme Sensors GmbH",
    angle: str = "Official product evidence. hardware signals: industrial, sensor, controller.",
    source: str = "https://acme.example/contact",
    industry: str = "Industrial Sensor Controllers",
    email: str = "sales@acme.example",
    email_type: str = "Generic Company Email",
    company_domain: str = "acme.example",
) -> dict[str, str]:
    body = " ".join(["message"] * 125) + "\n\n" + OPT_OUT_SENTENCE
    return {
        "Company Name": company,
        "Contact Email": email,
        "Email Type": email_type,
        "Subject": "Acme sensor controller projects",
        "Email Body": body,
        "Why This Angle Fits": angle,
        "Source URL": source,
        "Compliance Note": "Public business contact collected from the official company website.",
        "Company Domain": company_domain,
        "Country": "Germany",
        "Industry": industry,
        "Lead Score": "4",
    }


@pytest.mark.parametrize(
    ("company", "angle", "source", "industry"),
    [
        ("Global Travel Reviews", "Travel blog affiliate guide. hardware signals: sensor, controller.", "https://travel.example", "Travel"),
        ("Industry Directory", "Directory listing. hardware signals: industrial, sensor, controller.", "https://www.industrystock.com/en/company/acme", "Industrial Automation"),
        ("Contact", "Official product evidence. hardware signals: industrial, sensor, controller.", "https://acme.example/contact", "Industrial Automation"),
        ("Acme Software", "Official product evidence. hardware signals: sensor.", "https://acme.example/contact", "Software"),
    ],
)
def test_send_time_qualification_rejects_unqualified_drafts(manager, company, angle, source, industry):
    manager.append_outreach_drafts([sendable_draft(company, angle, source, industry)])
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"


def test_send_time_qualification_accepts_specific_recorded_discovery_hardware_signals(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                angle=(
                    "Attributable product evidence. hardware signals: sensor, controller. "
                    "Stock service angle offers embedded electronics support."
                )
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "YES"


def test_send_time_qualification_requires_eligible_hardware_evidence(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                angle=(
                    "Attributable service evidence. hardware signals: industrial, automation. "
                    "Stock service angle offers embedded electronics support."
                ),
                industry="Industrial Automation Consulting",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "Hardware qualification failed" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


def test_send_time_exclusion_replay_uses_full_recorded_angle_evidence(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                angle=(
                    "Official identity: SaaS software platform and consulting services. "
                    "hardware signals: industrial, sensor, controller."
                )
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    reason = ws.cell(2, headers.index("Auto Send Reason") + 1).value
    assert "software-only" in reason
    assert "consulting-only" in reason


def test_send_time_hardware_qualification_never_uses_template_body(manager):
    candidate = sendable_draft(angle="Official evidence. hardware signals: industrial.")
    candidate["Email Body"] = " ".join(["sensor controller hardware"] * 42) + "\n\n" + OPT_OUT_SENTENCE

    manager.append_outreach_drafts([candidate])
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "Hardware qualification failed" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


@pytest.mark.parametrize(
    "overrides",
    [
        {"source": "https://acme.example/online-magazine/about"},
        {"industry": "Tourism Site"},
        {"angle": "Featured by a travel magazine. hardware signals: industrial, sensor, controller."},
        {"angle": "Featured by a travel publisher. hardware signals: industrial, sensor, controller."},
    ],
)
def test_send_time_travel_and_media_terms_without_prospect_identity_remain_eligible(manager, overrides):
    manager.append_outreach_drafts([sendable_draft(**overrides)])
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "YES"


@pytest.mark.parametrize(
    ("company", "angle", "source", "expected_exclusion"),
    [
        (
            "Acme Software Developer",
            "We manufacture industrial sensor controllers. hardware signals: industrial, sensor, controller.",
            "https://acme.example/contact",
            "software-only",
        ),
        (
            "Acme Cloud",
            "We manufacture industrial sensor controllers. hardware signals: industrial, sensor, controller. "
            "Our company operates",
            "a cloud platform as its business; https://acme.example/contact",
            "software-only",
        ),
        (
            "Atlas Controls",
            "We manufacture industrial sensor controllers. hardware signals: industrial, sensor, controller. "
            "Our company operates",
            "a travel blog as its business; https://acme.example/contact",
            "travel/tourism",
        ),
        (
            "Technology Magazine",
            "We manufacture industrial sensor controllers. hardware signals: industrial, sensor, controller.",
            "https://acme.example/contact",
            "news/media",
        ),
    ],
)
def test_send_time_replays_company_identity_with_full_angle_and_source_evidence(
    manager,
    company,
    angle,
    source,
    expected_exclusion,
):
    manager.append_outreach_drafts([sendable_draft(company=company, angle=angle, source=source)])
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert expected_exclusion in ws.cell(2, headers.index("Auto Send Reason") + 1).value


@pytest.mark.parametrize(
    "identity_text",
    [
        "We sell our industrial sensor controller analytics as a cloud platform",
        "We offer our industrial sensor controller analytics as a cloud platform",
    ],
)
def test_send_time_rejects_cloud_platform_operator_identity_despite_physical_action(manager, identity_text):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                angle=f"{identity_text}. hardware signals: industrial, sensor, controller.",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "software-only" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


def test_send_time_rejects_first_party_travel_magazine_operator_identity(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                angle=(
                    "We are a travel magazine covering industrial sensor controllers. "
                    "hardware signals: industrial, sensor, controller."
                ),
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "travel/tourism" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


@pytest.mark.parametrize(
    ("email", "email_type", "company_domain", "source", "expected_reason"),
    [
        (
            "founder@acme.example",
            "Private Email",
            "acme.example",
            "https://acme.example/contact",
            "Email Type not allowed",
        ),
        (
            "sales@yahoo.com",
            "Generic Company Email",
            "acme.example",
            "https://acme.example/contact",
            "Free email domain not allowed",
        ),
        (
            "sales@unrelated.example",
            "Public Work Email",
            "acme.example",
            "https://acme.example/contact",
            "Recipient domain not attributable",
        ),
    ],
)
def test_send_time_eligibility_rejects_private_free_and_unrelated_recipients(
    manager,
    email,
    email_type,
    company_domain,
    source,
    expected_reason,
):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email=email,
                email_type=email_type,
                company_domain=company_domain,
                source=source,
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert expected_reason in ws.cell(2, headers.index("Auto Send Reason") + 1).value


@pytest.mark.parametrize(
    ("email", "email_type", "company_domain", "source"),
    [
        (
            "sales@acme.example",
            "Generic Company Email",
            "acme.example",
            "https://acme.example/contact",
        ),
        (
            "engineer@acme.example",
            "Public Work Email",
            "shop.acme.example",
            "https://shop.acme.example/about",
        ),
        (
            "sales@acme.example",
            "Generic Company Email",
            "",
            "https://www.acme.example/legal",
        ),
    ],
)
def test_send_time_eligibility_accepts_only_attributable_company_recipients(
    manager,
    email,
    email_type,
    company_domain,
    source,
):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email=email,
                email_type=email_type,
                company_domain=company_domain,
                source=source,
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "YES"


@pytest.mark.parametrize(
    "email",
    [
        "sales@@acme.co.uk",
        ".sales@acme.co.uk",
        "sales.@acme.co.uk",
        "sa..les@acme.co.uk",
        "sales@-acme.co.uk",
        "sales@acme..co.uk",
        "sales@acme.co.uk:25",
        "sales@co.uk",
    ],
)
def test_send_time_eligibility_uses_strict_full_mailbox_validation(manager, email):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email=email,
                company_domain="acme.co.uk",
                source="https://acme.co.uk/contact",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "Invalid email format" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


@pytest.mark.parametrize(
    "email",
    [
        "\u017fales@acme.co.uk",
        "sales@acme\u2024co.uk",
        "sales@acme\ufe52co.uk",
        "sales@b\u00fccher.de",
        "sales@.acme.co.uk",
        "sales@acme.co.uk.",
        "sales@www.acme.co.uk",
        "sales@acme-.co.uk",
        "sales@ac_me.co.uk",
    ],
)
def test_send_time_eligibility_rejects_lookalikes_and_raw_domain_mutations(manager, email):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email=email,
                company_domain="acme.co.uk",
                source="https://acme.co.uk/contact",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "Invalid email format" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


def test_send_time_eligibility_accepts_valid_dot_atom_mailbox_on_parent_company_domain(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email="first.last+sales@acme.co.uk",
                email_type="Public Work Email",
                company_domain="shop.acme.co.uk",
                source="https://shop.acme.co.uk/contact",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "YES"


def test_send_time_eligibility_rejects_child_mail_host_for_parent_official_host(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email="sales@mail.acme.co.uk",
                company_domain="acme.co.uk",
                source="https://acme.co.uk/contact",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "Recipient domain not attributable" in ws.cell(2, headers.index("Auto Send Reason") + 1).value


def test_send_time_eligibility_rejects_sibling_mail_host_for_official_subdomain(manager):
    manager.append_outreach_drafts(
        [
            sendable_draft(
                email="sales@mail.acme.co.uk",
                company_domain="shop.acme.co.uk",
                source="https://shop.acme.co.uk/contact",
            )
        ]
    )
    ws = manager.sheet(OUTREACH_SHEET)
    headers = manager.headers(ws)

    assert ws.cell(2, headers.index("Auto Send Eligible") + 1).value == "NO"
    assert "Recipient domain not attributable" in ws.cell(2, headers.index("Auto Send Reason") + 1).value
