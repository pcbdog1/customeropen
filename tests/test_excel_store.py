from pcb_leads.excel_store import ExcelLeadStore
from pcb_leads.models import Lead


def make_lead(company: str, website: str) -> Lead:
    return Lead(
        country="Germany",
        city="Berlin",
        company_name=company,
        website=website,
        business="Robotics hardware",
        industry="Robotics",
        pcb_need_reason="Robot controllers contain electronics and PCBAs.",
        demand_signal="Hardware products",
        demand_signal_source="Company website",
        employees="未找到",
        founded="未找到",
        email="info@example.com",
        email_type="Generic Company Email",
        contact_method="Email",
        contact_form_url="",
        contact_name="未公开/待核实",
        contact_title="未公开/待核实",
        phone="未公开/待核实",
        linkedin="未公开/待核实",
        source_links="https://example.com",
        score=4,
        development_angle="机器人控制板PCBA打样",
        email_subject="PCB & PCBA Support for Robotics Control Electronics",
        notes="",
        compliance_note="Public business contact collected from source URL. B2B relevance: PCB/PCBA manufacturing service matches company hardware/electronics business. Outreach email must include opt-out sentence.",
    )


def test_excel_store_appends_and_skips_duplicate_company_domain(tmp_path):
    path = tmp_path / "leads.xlsx"
    store = ExcelLeadStore(path)

    result = store.append_leads(
        [
            make_lead("ACME Robotics GmbH", "https://www.acme.de"),
            make_lead(" ACME Robotics GmbH ", "https://acme.de/products"),
        ]
    )

    assert result.added == 1
    assert result.skipped_duplicates == 1

    reloaded = ExcelLeadStore(path)
    assert len(reloaded.load_existing_keys()) == 1


def test_excel_store_writes_new_email_and_compliance_columns(tmp_path):
    path = tmp_path / "leads.xlsx"
    ExcelLeadStore(path).append_leads([make_lead("ACME Robotics GmbH", "https://www.acme.de")])

    from openpyxl import load_workbook

    ws = load_workbook(path)["Leads"]
    headers = [cell.value for cell in ws[1]]
    assert "Email Type" in headers
    assert "Contact Method" in headers
    assert "Contact Form URL" in headers
    assert "Compliance Note" in headers
