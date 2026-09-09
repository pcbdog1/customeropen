from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from itertools import chain
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from excel_manager import ExcelManager, LEADS_SHEET, OUTREACH_SHEET
from pcb_leads.models import Lead
from pcb_leads.parser import FREE_EMAIL_DOMAINS
from pcb_leads.qualification import is_generic_company_name
from pcb_leads.utils import company_country_key, normalize_host, normalize_mailbox, registrable_domain


COMPLIANCE_NOTE = "Imported from user-provided verified B2B customer email list for PCB/PCBA outreach."
HARDWARE_SIGNAL = "User-provided PCB/PCBA potential customer list"
RECORDED_HARDWARE_SIGNAL = f"Hardware signals: electronics; embedded. {HARDWARE_SIGNAL}"
EMAIL_TYPE = "User Provided Verified Business Email"
EMAIL_RE = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.IGNORECASE)

FIELD_ALIASES = {
    "company_name": {
        "company name", "company", "name", "公司名称", "公司名称(en)", "公司名称（en）",
        "展商英文名称", "展商中文名称",
    },
    "website": {"website", "company website", "company url", "url", "网站", "官网", "公司官网"},
    "domain": {"domain", "company domain", "域名", "官网域名"},
    "contact_person": {"contact person", "contact name", "contact", "联系人", "联系人1", "联系人2"},
    "country": {"country", "国家", "国家/地区", "国家／地区"},
    "city": {"city", "城市", "市", "城市/国家"},
    "industry": {
        "industry", "行业", "客户领域", "命中领域(全部)", "命中领域（全部）", "应用领域",
        "产品服务", "产品/服务(摘要)", "产品/服务（摘要）", "产品类别1",
    },
    "source_url": {"source url", "source", "source link", "来源url", "来源链接", "数据来源链接"},
    "notes": {
        "notes", "note", "remarks", "备注", "判定依据", "展商优势", "展商简介", "公司简介(摘要)",
        "公司简介（摘要）",
    },
}
EMAIL_HEADERS = {
    "email", "email address", "contact email", "e-mail", "邮箱", "邮箱1", "邮箱2", "邮箱3",
    "企业邮箱", "企业邮箱1", "企业邮箱2", "企业邮箱3", "联系人邮箱",
}


@dataclass
class EmailSeedImportStats:
    files_read: int = 0
    raw_emails: int = 0
    skipped_duplicates: int = 0
    skipped_sent: int = 0
    new_customers: int = 0
    new_drafts: int = 0
    pending_eligible: int = 0


def _normalise_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _field_map(first_row: tuple[object, ...]) -> dict[str, list[int]]:
    result: dict[str, list[int]] = {}
    for index, value in enumerate(first_row):
        header = _normalise_header(value)
        if header in EMAIL_HEADERS or header.startswith("邮箱") or header.startswith("企业邮箱"):
            result.setdefault("email", []).append(index)
            continue
        for field, aliases in FIELD_ALIASES.items():
            if header in aliases:
                result.setdefault(field, []).append(index)
                break
    return result


def _xlsx_tables(path: Path) -> Iterator[tuple[dict[str, list[int]], Iterable[tuple[object, ...]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    for worksheet in workbook.worksheets:
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, tuple())
        mapping = _field_map(first)
        if "email" in mapping:
            yield mapping, iterator
        elif any(EMAIL_RE.search(str(value or "")) for value in first):
            yield {"email": list(range(len(first)))}, chain((first,), iterator)


def _xls_tables(path: Path) -> Iterator[tuple[dict[str, list[int]], Iterable[tuple[object, ...]]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("Reading .xls seed files requires xlrd. Run: pip install -r requirements.txt") from exc

    workbook = xlrd.open_workbook(path)
    for worksheet in workbook.sheets():
        if not worksheet.nrows:
            continue
        first = tuple(worksheet.row_values(0))
        mapping = _field_map(first)
        rows = (tuple(worksheet.row_values(row)) for row in range(1, worksheet.nrows))
        if "email" in mapping:
            yield mapping, rows
        elif any(EMAIL_RE.search(str(value or "")) for value in first):
            yield {"email": list(range(len(first)))}, chain((first,), rows)


def _csv_tables(path: Path) -> Iterator[tuple[dict[str, list[int]], Iterable[tuple[object, ...]]]]:
    handle = path.open("r", encoding="utf-8-sig", newline="")
    reader = csv.reader(handle)
    first = tuple(next(reader, []))
    mapping = _field_map(first)
    if "email" in mapping:
        yield mapping, reader
    else:
        yield {"email": list(range(len(first)))}, chain((first,), reader)


def _tables(path: Path) -> Iterator[tuple[dict[str, list[int]], Iterable[tuple[object, ...]]]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        yield from _csv_tables(path)
    elif suffix == ".xls":
        yield from _xls_tables(path)
    else:
        yield from _xlsx_tables(path)


def _value(row: tuple[object, ...], mapping: dict[str, list[int]], field: str) -> str:
    for index in mapping.get(field, []):
        if index < len(row) and str(row[index] or "").strip():
            return str(row[index]).strip()
    return ""


def _emails(row: tuple[object, ...], mapping: dict[str, list[int]]) -> Iterator[str]:
    for index in mapping.get("email", []):
        if index >= len(row):
            continue
        for match in EMAIL_RE.findall(str(row[index] or "")):
            email = normalize_mailbox(match)
            if email:
                yield email


def _company_from_domain(domain: str) -> str:
    label = domain.split(".", 1)[0].replace("-", " ").replace("_", " ").strip()
    return label.title() or domain


def _domain_key(value: object) -> str:
    host = normalize_host(value)
    return registrable_domain(host) or host


def _known_domains(manager: ExcelManager, sheet_name: str) -> set[str]:
    worksheet = manager.sheet(sheet_name)
    headers = manager.headers(worksheet)
    candidates = ("官网域名", "公司官网", "Company Domain", "Source URL")
    domains: set[str] = set()
    for column_name in candidates:
        if column_name not in headers:
            continue
        column = headers.index(column_name) + 1
        domains.update(
            domain for row in range(2, worksheet.max_row + 1)
            if (domain := _domain_key(worksheet.cell(row, column).value))
        )
    return domains


def _known_company_countries(manager: ExcelManager, sheet_name: str) -> set[str]:
    worksheet = manager.sheet(sheet_name)
    headers = manager.headers(worksheet)
    company_header = "公司名称" if "公司名称" in headers else "Company Name"
    country_header = "国家" if "国家" in headers else "Country"
    if company_header not in headers or country_header not in headers:
        return set()
    company_column = headers.index(company_header) + 1
    country_column = headers.index(country_header) + 1
    keys: set[str] = set()
    for row in range(2, worksheet.max_row + 1):
        company = str(worksheet.cell(row, company_column).value or "").strip()
        country = str(worksheet.cell(row, country_column).value or "").strip()
        if company and country and not is_generic_company_name(company):
            keys.add(company_country_key(company, country))
    return keys


def import_email_seed_files(
    manager: ExcelManager,
    sources: list[Path],
    draft_builder: Callable[[Lead], dict[str, Any]],
) -> EmailSeedImportStats:
    sent_emails = manager.sent_emails()
    sent_domains = {_domain_key(domain) for domain in manager.sent_domains() if _domain_key(domain)}
    sent_companies = manager.sent_companies()
    existing_emails = manager.existing_emails()
    existing_domains = _known_domains(manager, LEADS_SHEET) | _known_domains(manager, OUTREACH_SHEET)
    existing_company_countries = (
        _known_company_countries(manager, LEADS_SHEET)
        | _known_company_countries(manager, OUTREACH_SHEET)
    )
    seen_emails: set[str] = set()
    seen_domains: set[str] = set()
    seen_company_countries: set[str] = set()
    leads: list[Lead] = []
    stats = EmailSeedImportStats(files_read=len(sources))

    for source in sources:
        for mapping, rows in _tables(source):
            for raw_row in rows:
                row = tuple(raw_row)
                for email in _emails(row, mapping):
                    stats.raw_emails += 1
                    email_domain = _domain_key(email.partition("@")[2])
                    supplied_domain = _domain_key(_value(row, mapping, "domain")) or _domain_key(_value(row, mapping, "website"))
                    domain = supplied_domain if supplied_domain == email_domain else email_domain
                    company = _value(row, mapping, "company_name") or _company_from_domain(domain)
                    country = _value(row, mapping, "country") or "Not Found"
                    company_country = company_country_key(company, country)
                    if email in sent_emails or domain in sent_domains or company.casefold() in sent_companies:
                        stats.skipped_sent += 1
                        continue
                    if (
                        domain in FREE_EMAIL_DOMAINS
                        or email in existing_emails
                        or domain in existing_domains
                        or (not is_generic_company_name(company) and company_country in existing_company_countries)
                        or email in seen_emails
                        or domain in seen_domains
                        or (not is_generic_company_name(company) and company_country in seen_company_countries)
                    ):
                        stats.skipped_duplicates += 1
                        continue

                    supplied_website = _value(row, mapping, "website").rstrip("/")
                    website = supplied_website if _domain_key(supplied_website) == domain else f"https://{domain}"
                    if not website.startswith(("http://", "https://")):
                        website = f"https://{website.lstrip('/')}"
                    supplied_source = _value(row, mapping, "source_url")
                    source_url = supplied_source if _domain_key(supplied_source) == domain else website
                    industry = _value(row, mapping, "industry") or "Hardware / Electronics"
                    notes = _value(row, mapping, "notes") or f"Imported from {source.name}."
                    contact_name = _value(row, mapping, "contact_person") or "Not Found"
                    leads.append(Lead(
                        country=country,
                        city=_value(row, mapping, "city") or "Not Found",
                        company_name=company,
                        website=website,
                        business=HARDWARE_SIGNAL,
                        industry=industry,
                        pcb_need_reason=RECORDED_HARDWARE_SIGNAL,
                        demand_signal=HARDWARE_SIGNAL,
                        demand_signal_source=source_url,
                        employees="Not Found",
                        founded="Not Found",
                        email=email,
                        email_type=EMAIL_TYPE,
                        contact_method="User-provided verified business email",
                        contact_form_url="",
                        contact_name=contact_name,
                        contact_title="Not Found",
                        phone="Not Found",
                        linkedin="",
                        source_links=source_url,
                        score=4,
                        development_angle=RECORDED_HARDWARE_SIGNAL,
                        email_subject="PCB/PCBA support for your hardware programs",
                        notes=notes,
                        compliance_note=COMPLIANCE_NOTE,
                    ))
                    seen_emails.add(email)
                    seen_domains.add(domain)
                    if not is_generic_company_name(company):
                        seen_company_countries.add(company_country)

    appended = manager.append_leads(leads)
    drafts: list[dict[str, Any]] = []
    for lead in leads:
        draft = draft_builder(lead)
        draft["Email Source URL"] = lead.source_links
        drafts.append(draft)
    stats.new_customers = appended.added
    stats.skipped_duplicates += appended.skipped_duplicates
    stats.new_drafts = manager.append_outreach_drafts(drafts)
    manager.evaluate_auto_send_all()
    stats.pending_eligible = manager.pending_auto_eligible_count()
    return stats


def import_email_seeds(
    manager: ExcelManager,
    source: Path,
    draft_builder: Callable[[Lead], dict[str, Any]],
) -> EmailSeedImportStats:
    return import_email_seed_files(manager, [source], draft_builder)
