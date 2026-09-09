from __future__ import annotations

import os
import random
import smtplib
import subprocess
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from email.message import EmailMessage

from openpyxl.worksheet.worksheet import Worksheet

from email_sender_common import OPT_OUT_SENTENCE
from excel_manager import ExcelManager, OUTREACH_SHEET
from pcb_leads.utils import normalize_host


@dataclass
class SendableEmail:
    row_number: int
    company: str
    email: str
    email_type: str
    subject: str
    body: str
    country: str
    industry: str
    source_url: str
    domain: str


@dataclass
class SendStats:
    selected: int = 0
    attempted: int = 0
    sent: int = 0
    failed: int = 0
    skipped: int = 0
    throttle_reason: str = ""
    last_customer: str = ""
    last_email: str = ""
    sent_items: list[SendableEmail] = field(default_factory=list)

    @property
    def failure_rate(self) -> float:
        return self.failed / self.attempted if self.attempted else 0.0


def _header_map(ws: Worksheet) -> dict[str, int]:
    return {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}


def select_auto_send_rows(manager: ExcelManager, max_send: int) -> list[SendableEmail]:
    ws = manager.sheet(OUTREACH_SHEET)
    manager.evaluate_auto_send_all()
    headers = _header_map(ws)
    selected: list[SendableEmail] = []
    sent_emails = manager.sent_emails()
    sent_companies = manager.sent_companies()
    sent_domains = manager.sent_domains()
    domains_sent_today: set[str] = set()
    today = date.today().isoformat()
    log = manager.sheet("Sent_Log")
    log_headers = _header_map(log)
    if "Company Domain" in log_headers and "Send Date" in log_headers and "Status" in log_headers:
        for row in range(2, log.max_row + 1):
            send_date = str(log.cell(row, log_headers["Send Date"]).value or "")
            status = str(log.cell(row, log_headers["Status"]).value or "")
            domain = str(log.cell(row, log_headers["Company Domain"]).value or "")
            if send_date.startswith(today) and status == "SENT" and domain:
                domains_sent_today.add(domain)

    for row in range(2, ws.max_row + 1):
        if len(selected) >= max_send:
            break
        eligible = str(ws.cell(row, headers["Auto Send Eligible"]).value or "").upper()
        status = str(ws.cell(row, headers["Send Status"]).value or "")
        if eligible != "YES" or status:
            continue
        final_safe, _ = manager.final_send_safety_check(ws, headers, row)
        if not final_safe:
            continue
        email = str(ws.cell(row, headers["Contact Email"]).value or "").strip()
        company = str(ws.cell(row, headers["Company Name"]).value or "").strip()
        domain = normalize_host(ws.cell(row, headers["Company Domain"]).value)
        if email.lower() in sent_emails or company.lower() in sent_companies or domain in sent_domains or domain in domains_sent_today:
            continue
        body = str(ws.cell(row, headers["Email Body"]).value or "")
        if OPT_OUT_SENTENCE not in body:
            continue
        selected.append(
            SendableEmail(
                row_number=row,
                company=company,
                email=email,
                email_type=str(ws.cell(row, headers["Email Type"]).value or ""),
                subject=str(ws.cell(row, headers["Subject"]).value or ""),
                body=body,
                country=str(ws.cell(row, headers["Country"]).value or ""),
                industry=str(ws.cell(row, headers["Industry"]).value or ""),
                source_url=str(ws.cell(row, headers["Source URL"]).value or ""),
                domain=domain,
            )
        )
        sent_emails.add(email.lower())
        sent_companies.add(company.lower())
        domains_sent_today.add(domain)
    return selected


class RealEmailSender:
    def __init__(self):
        self.smtp_host = os.getenv("SMTP_HOST", "")
        self.smtp_port = int(os.getenv("SMTP_PORT", "587"))
        self.smtp_user = os.getenv("SMTP_USER", "")
        self.smtp_password = os.getenv("SMTP_PASSWORD", "")
        self.smtp_from = os.getenv("SMTP_FROM", self.smtp_user)
        self.mail_account = os.getenv("MAIL_APP_ACCOUNT", "")
        self.mail_sender = os.getenv("MAIL_APP_SENDER", self.smtp_from)

    def send(self, item: SendableEmail) -> None:
        if self.smtp_host and self.smtp_user and self.smtp_password and self.smtp_from:
            self._send_smtp(item)
        else:
            self._send_mail_app(item)

    def _send_smtp(self, item: SendableEmail) -> None:
        message = EmailMessage()
        message["From"] = self.smtp_from
        message["To"] = item.email
        message["Subject"] = item.subject
        message.set_content(item.body)
        with smtplib.SMTP(self.smtp_host, self.smtp_port, timeout=30) as smtp:
            smtp.starttls()
            smtp.login(self.smtp_user, self.smtp_password)
            smtp.send_message(message)

    def _send_mail_app(self, item: SendableEmail) -> None:
        script = """
on run argv
    set recipientAddress to item 1 of argv
    set messageSubject to item 2 of argv
    set messageBody to item 3 of argv
    tell application "Mail"
        set newMessage to make new outgoing message with properties {subject:messageSubject, content:messageBody, visible:false}
        tell newMessage
            set sender to "%s"
            make new to recipient at end of to recipients with properties {address:recipientAddress}
            send
        end tell
    end tell
end run
""" % self.mail_sender
        completed = subprocess.run(
            ["osascript", "-", item.email, item.subject, item.body],
            input=script,
            text=True,
            capture_output=True,
            timeout=60,
        )
        if completed.returncode != 0:
            raise RuntimeError((completed.stderr or completed.stdout or "Mail.app send failed").strip())


def send_auto_eligible_emails(manager: ExcelManager, max_send: int, wait_min: int, wait_max: int, dry_run: bool = False) -> SendStats:
    ws = manager.sheet(OUTREACH_SHEET)
    headers = _header_map(ws)
    selected = select_auto_send_rows(manager, max_send)
    stats = SendStats(selected=len(selected))
    if dry_run:
        return stats
    sender = RealEmailSender()
    stats.attempted = len(selected)
    for idx, item in enumerate(selected):
        now = datetime.now().isoformat(timespec="seconds")
        try:
            manager.evaluate_auto_send_all()
            sent_emails = manager.sent_emails()
            sent_companies = manager.sent_companies()
            sent_domains = manager.sent_domains()
            current_status = str(ws.cell(item.row_number, headers["Send Status"]).value or "").strip()
            current_eligible = str(ws.cell(item.row_number, headers["Auto Send Eligible"]).value or "").upper()
            final_safe, _ = manager.final_send_safety_check(ws, headers, item.row_number)
            if (
                current_status
                or current_eligible != "YES"
                or not final_safe
                or item.email.lower() in sent_emails
                or item.company.lower() in sent_companies
                or normalize_host(item.domain) in sent_domains
            ):
                stats.skipped += 1
                continue
            sender.send(item)
            ws.cell(item.row_number, headers["Send Status"], "SENT")
            ws.cell(item.row_number, headers["Send Date"], now)
            ws.cell(item.row_number, headers["Error Message"], "")
            manager.append_sent_log(item.company, item.email, item.subject, "SENT", "", item.country, item.industry, item.source_url, item.domain)
            manager.save()
            stats.last_customer = item.company
            stats.last_email = item.email
            stats.sent_items.append(item)
            stats.sent += 1
        except Exception as exc:
            ws.cell(item.row_number, headers["Send Status"], "FAILED")
            ws.cell(item.row_number, headers["Error Message"], str(exc))
            manager.append_sent_log(item.company, item.email, item.subject, "FAILED", str(exc), item.country, item.industry, item.source_url, item.domain)
            stats.failed += 1
        if idx < len(selected) - 1:
            time.sleep(random.randint(wait_min, wait_max))
    if stats.failed > 5:
        stats.throttle_reason = "Today failed sends exceeded 5; next run should throttle to 20."
    return stats
