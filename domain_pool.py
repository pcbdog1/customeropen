from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

from pcb_leads.utils import normalize_domain, normalize_host


POOL_HEADERS = [
    "Date Added", "Company Name", "Domain", "Website", "Country", "Industry",
    "Source Type", "Source URL", "Hardware Signal", "Email Search Status",
    "Email Found", "Last Checked Date", "Send Status", "Notes",
]


@dataclass(frozen=True)
class DomainPoolRecord:
    company_name: str
    domain: str
    website: str
    country: str
    industry: str
    source_type: str
    source_url: str
    hardware_signal: str


class DomainPool:
    """Append-only local domain pool used before any optional external API."""

    def __init__(self, path: Path = Path("outputs/company_domain_pool.xlsx")) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not path.exists():
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Company_Domain_Pool"
            worksheet.append(POOL_HEADERS)
            workbook.save(path)
        self.workbook = load_workbook(path)
        self.worksheet = self.workbook["Company_Domain_Pool"]
        self._ensure_headers()

    def _ensure_headers(self) -> None:
        headers = [cell.value for cell in self.worksheet[1]]
        for header in POOL_HEADERS:
            if header not in headers:
                self.worksheet.cell(1, len(headers) + 1, header)
                headers.append(header)

    @property
    def headers(self) -> dict[str, int]:
        return {value: index + 1 for index, value in enumerate(cell.value for cell in self.worksheet[1]) if value}

    def _existing_domains(self) -> set[str]:
        domain_column = self.headers["Domain"]
        return {
            normalize_domain(value)
            for (value,) in self.worksheet.iter_rows(min_row=2, min_col=domain_column, max_col=domain_column, values_only=True)
            if normalize_domain(value)
        }

    def add(self, records: Iterable[DomainPoolRecord]) -> int:
        existing = self._existing_domains()
        added = 0
        for record in records:
            domain = normalize_domain(record.domain or record.website)
            if not domain or domain in existing:
                continue
            website = record.website.strip() if record.website else f"https://{domain}"
            self.worksheet.append([
                date.today().isoformat(), record.company_name.strip() or domain, domain, website,
                record.country, record.industry, record.source_type, record.source_url or website,
                record.hardware_signal, "Unchecked", "", "", "", "",
            ])
            existing.add(domain)
            added += 1
        return added

    def unchecked(self, limit: int) -> list[DomainPoolRecord]:
        h = self.headers
        records: list[DomainPoolRecord] = []
        for row in range(2, self.worksheet.max_row + 1):
            if str(self.worksheet.cell(row, h["Email Search Status"]).value or "").strip().casefold() != "unchecked":
                continue
            domain = normalize_domain(self.worksheet.cell(row, h["Domain"]).value)
            if not domain:
                continue
            records.append(DomainPoolRecord(
                company_name=str(self.worksheet.cell(row, h["Company Name"]).value or domain),
                domain=domain,
                website=str(self.worksheet.cell(row, h["Website"]).value or f"https://{domain}"),
                country=str(self.worksheet.cell(row, h["Country"]).value or ""),
                industry=str(self.worksheet.cell(row, h["Industry"]).value or ""),
                source_type=str(self.worksheet.cell(row, h["Source Type"]).value or "Local Seed"),
                source_url=str(self.worksheet.cell(row, h["Source URL"]).value or ""),
                hardware_signal=str(self.worksheet.cell(row, h["Hardware Signal"]).value or ""),
            ))
            if len(records) >= limit:
                break
        return records

    def mark_checked(self, domains: Iterable[str], emails_by_domain: dict[str, str] | None = None) -> None:
        h = self.headers
        wanted = {normalize_domain(domain) for domain in domains if normalize_domain(domain)}
        found = {normalize_domain(domain): email for domain, email in (emails_by_domain or {}).items()}
        for row in range(2, self.worksheet.max_row + 1):
            domain = normalize_domain(self.worksheet.cell(row, h["Domain"]).value)
            if domain not in wanted:
                continue
            email = found.get(domain, "")
            self.worksheet.cell(row, h["Email Search Status"], "Found" if email else "Not Found")
            self.worksheet.cell(row, h["Email Found"], email)
            self.worksheet.cell(row, h["Last Checked Date"], date.today().isoformat())

    def mark_known(self, domains: Iterable[str]) -> int:
        """Remove historical leads from the unchecked queue without deleting audit data."""
        h = self.headers
        known = {normalize_domain(domain) for domain in domains if normalize_domain(domain)}
        updated = 0
        for row in range(2, self.worksheet.max_row + 1):
            if normalize_domain(self.worksheet.cell(row, h["Domain"]).value) not in known:
                continue
            if str(self.worksheet.cell(row, h["Email Search Status"]).value or "").strip().casefold() == "unchecked":
                self.worksheet.cell(row, h["Email Search Status"], "Already Known")
                self.worksheet.cell(row, h["Notes"], "Already present in lead workbook or sent history.")
                updated += 1
        return updated

    def metrics(self) -> dict[str, int]:
        h = self.headers
        total = max(0, self.worksheet.max_row - 1)
        unchecked = found = 0
        for row in range(2, self.worksheet.max_row + 1):
            status = str(self.worksheet.cell(row, h["Email Search Status"]).value or "").strip().casefold()
            if status == "unchecked":
                unchecked += 1
            if status == "found" or str(self.worksheet.cell(row, h["Email Found"]).value or "").strip():
                found += 1
        return {"total": total, "unchecked": unchecked, "email_found": found}

    def save(self) -> None:
        self.workbook.save(self.path)


def import_seed_file(pool: DomainPool, path: Path) -> int:
    if not path.exists():
        return 0
    rows: list[dict[str, object]] = []
    if path.suffix.casefold() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream))
    elif path.suffix.casefold() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    else:
        return 0
    def value(row: dict[str, object], *names: str) -> str:
        lookup = {str(key).casefold(): val for key, val in row.items()}
        for name in names:
            item = lookup.get(name.casefold())
            if item:
                return str(item).strip()
        return ""
    return pool.add(DomainPoolRecord(
        company_name=value(row, "Company Name", "Company", "Name"),
        domain=value(row, "Domain", "Website", "URL"),
        website=value(row, "Website", "URL") or f"https://{normalize_host(value(row, 'Domain'))}",
        country=value(row, "Country"), industry=value(row, "Industry"),
        source_type=value(row, "Source Type") or "Imported Seed",
        source_url=value(row, "Source URL", "Website", "URL"),
        hardware_signal=value(row, "Hardware Signal", "Notes"),
    ) for row in rows)
