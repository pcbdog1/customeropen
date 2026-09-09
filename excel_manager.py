from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from email_sender_common import OPT_OUT_SENTENCE
from pcb_leads.directory_sources import NON_COMPANY_DISCOVERY_DOMAINS, is_directory_url
from pcb_leads.excel_store import HEADERS as LEAD_HEADERS, LEADS_SHEET
from pcb_leads.models import Lead
from pcb_leads.parser import FREE_EMAIL_DOMAINS
from pcb_leads.qualification import hardware_qualification, is_generic_company_name
from pcb_leads.scorer import is_likely_peer_manufacturer
from pcb_leads.utils import (
    company_country_key,
    dedupe_key,
    hosts_are_attributable,
    mailbox_domain,
    normalize_company_name,
    normalize_domain,
    normalize_host,
    normalize_mailbox,
    registrable_domain,
)

OUTREACH_SHEET = "Outreach_Email_Drafts"
SENT_LOG_SHEET = "Sent_Log"
DAILY_REPORT_SHEET = "Daily_Report"
EMAIL_RECOVERY_SHEET = "Email_Recovery_Report"

OUTREACH_HEADERS = [
    "Company Name",
    "Contact Email",
    "Email Type",
    "Subject",
    "Email Body",
    "Why This Angle Fits",
    "Auto Send Eligible",
    "Auto Send Reason",
    "Send Status",
    "Send Date",
    "Error Message",
    "Source URL",
    "Compliance Note",
    "Do Not Contact",
    "Lead Score",
    "Company Domain",
    "Country",
    "Industry",
    "Word Count",
    "Industry Emphasis",
    "Final Send Safety Check",
    "Final Send Block Reason",
    "Email Source URL",
]

_RECORDED_HARDWARE_SIGNALS_RE = re.compile(r"\bhardware signals?\s*:\s*([^\.\n]+)", re.IGNORECASE)
_UNRESOLVED_COMPANY_NAMES = {"", "not found", "unknown", "n/a", "na", "未找到"}
_ALLOWED_SEND_EMAIL_TYPES = {
    "Generic Company Email",
    "Public Work Email",
    "User Provided Verified Business Email",
}
_BLOCKED_SEND_HOSTS = frozenset({"wikipedia.org", "wikimedia.org"})
_BLOCKED_SOURCE_TERMS = (
    "category:", "category/", "/category/", "list of companies",
    "government", "university", "college", "blog", "news", "affiliate", "travel",
    "marketing agency", "seo agency", "recruitment", "staffing", "job board",
)
_BLOCKED_COMPANY_NAMES = frozenset({
    "wikipedia", "category", "list", "contact", "contact us", "legal notice", "disclaimer",
    "home", "about", "about us",
})
_PAGE_TITLE_COMPANY_NAME_RE = re.compile(
    r"(?:\bcategory\b|\blist\b|\bcompanies\b|\bdevices?\s+is\b|\bis\s+[A-Z])",
    re.IGNORECASE,
)
SENT_LOG_HEADERS = [
    "Send Date",
    "Company Name",
    "Contact Email",
    "Subject",
    "Status",
    "Error Message",
    "Country",
    "Industry",
    "Source URL",
    "Company Domain",
]
EMAIL_RECOVERY_HEADERS = [
    "Company Name",
    "Website",
    "Old Email",
    "New Email",
    "Email Type",
    "Email Source URL",
    "Recovery Method",
    "Safety Check Result",
    "Block Reason",
]


@dataclass
class AppendStats:
    added: int
    skipped_duplicates: int
    score4: int
    score5: int


class ExcelManager:
    def __init__(self, path: Path):
        self.path = path
        self.wb = load_workbook(path)
        self.ensure_schema()

    def ensure_schema(self) -> None:
        leads = self.wb[LEADS_SHEET]
        self._ensure_headers(leads, LEAD_HEADERS)
        outreach = self.wb[OUTREACH_SHEET] if OUTREACH_SHEET in self.wb.sheetnames else self.wb.create_sheet(OUTREACH_SHEET)
        self._delete_header_column(outreach, "Send Approval")
        self._rename_header(outreach, "Why this email angle fits this company", "Why This Angle Fits")
        self._ensure_headers(outreach, OUTREACH_HEADERS)
        self.backfill_outreach_metadata()
        sent_log = self.wb[SENT_LOG_SHEET] if SENT_LOG_SHEET in self.wb.sheetnames else self.wb.create_sheet(SENT_LOG_SHEET)
        self._ensure_headers(sent_log, SENT_LOG_HEADERS)
        if DAILY_REPORT_SHEET not in self.wb.sheetnames:
            self.wb.create_sheet(DAILY_REPORT_SHEET)
        recovery = self.wb[EMAIL_RECOVERY_SHEET] if EMAIL_RECOVERY_SHEET in self.wb.sheetnames else self.wb.create_sheet(EMAIL_RECOVERY_SHEET)
        self._ensure_headers(recovery, EMAIL_RECOVERY_HEADERS)
        self.evaluate_auto_send_all()

    @staticmethod
    def _ensure_headers(ws: Worksheet, headers: list[str]) -> None:
        existing = [cell.value for cell in ws[1]]
        if all(value is None for value in existing):
            for idx, header in enumerate(headers, start=1):
                ws.cell(row=1, column=idx, value=header)
            return
        for header in headers:
            if header not in existing:
                ws.cell(row=1, column=len(existing) + 1, value=header)
                existing.append(header)

    @staticmethod
    def _delete_header_column(ws: Worksheet, header: str) -> None:
        headers = [cell.value for cell in ws[1]]
        if header in headers:
            ws.delete_cols(headers.index(header) + 1)

    @staticmethod
    def _rename_header(ws: Worksheet, old: str, new: str) -> None:
        headers = [cell.value for cell in ws[1]]
        if old in headers and new not in headers:
            ws.cell(row=1, column=headers.index(old) + 1, value=new)

    @staticmethod
    def headers(ws: Worksheet) -> list[str]:
        return [cell.value for cell in ws[1]]

    def sheet(self, name: str) -> Worksheet:
        return self.wb[name]

    def lead_records(self) -> list[dict[str, Any]]:
        ws = self.wb[LEADS_SHEET]
        headers = self.headers(ws)
        return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]

    @staticmethod
    def _email_domain(email: Any) -> str:
        value = str(email or "").strip().casefold()
        return normalize_host(value.rsplit("@", 1)[1]) if "@" in value else ""

    @staticmethod
    def _valid_host(value: Any) -> str:
        host = normalize_host(str(value or ""))
        return host if "." in host else ""

    @classmethod
    def _valid_email_domain(cls, email: Any) -> str:
        return mailbox_domain(str(email or ""))

    @staticmethod
    def _hosts_are_attributable(first: str, second: str) -> bool:
        return hosts_are_attributable(first, second)

    @classmethod
    def _source_hosts(cls, source: Any) -> tuple[str, ...]:
        hosts: list[str] = []
        for raw_url in re.findall(r"https?://[^\s;]+", str(source or ""), flags=re.IGNORECASE):
            url = raw_url.rstrip(".,:)]}")
            parsed = urlparse(url)
            host = cls._valid_host(url)
            if (
                parsed.scheme.casefold() not in {"http", "https"}
                or not parsed.netloc
                or not host
                or normalize_domain(host) in NON_COMPANY_DISCOVERY_DOMAINS
                or is_directory_url(url)
            ):
                continue
            if host not in hosts:
                hosts.append(host)
        return tuple(hosts)

    @classmethod
    def _recipient_is_attributable(cls, email: Any, company_domain: Any, source: Any) -> bool:
        email_domain = cls._valid_email_domain(email)
        if not email_domain or email_domain in FREE_EMAIL_DOMAINS:
            return False
        company_host = cls._valid_host(company_domain)
        source_hosts = cls._source_hosts(source)
        if company_host:
            authoritative_hosts = [company_host]
            authoritative_hosts.extend(
                host for host in source_hosts if cls._hosts_are_attributable(host, company_host)
            )
        else:
            authoritative_hosts = list(source_hosts)
        return any(cls._hosts_are_attributable(email_domain, host) for host in authoritative_hosts)

    @classmethod
    def _lead_domains(cls, lead: dict[str, Any]) -> set[str]:
        return {
            domain
            for domain in (
                normalize_host(lead.get("公司官网")),
                cls._email_domain(lead.get("公开邮箱")),
            )
            if domain
        }

    @classmethod
    def _matching_lead(cls, outreach: dict[str, Any], leads: list[dict[str, Any]]) -> dict[str, Any]:
        outreach_domains = {
            domain
            for domain in (
                normalize_host(outreach.get("Company Domain")),
                cls._email_domain(outreach.get("Contact Email")),
            )
            if domain
        }
        if outreach_domains:
            matches = [lead for lead in leads if outreach_domains & cls._lead_domains(lead)]
            return matches[0] if len(matches) == 1 else {}

        company = str(outreach.get("Company Name") or "").strip()
        country = str(outreach.get("Country") or "").strip()
        if not company or not country or is_generic_company_name(company):
            return {}
        identity = company_country_key(company, country)
        matches = [
            lead
            for lead in leads
            if normalize_company_name(lead.get("公司名称"))
            and company_country_key(lead.get("公司名称"), lead.get("国家")) == identity
        ]
        return matches[0] if len(matches) == 1 else {}

    def backfill_outreach_metadata(self) -> None:
        ws = self.wb[OUTREACH_SHEET]
        headers = self.headers(ws)
        leads = self.lead_records()
        for row in range(2, ws.max_row + 1):
            outreach = {header: ws.cell(row, column).value for column, header in enumerate(headers, start=1)}
            lead = self._matching_lead(outreach, leads)
            mappings = {
                "Source URL": lead.get("数据来源链接", ""),
                "Compliance Note": lead.get("Compliance Note", ""),
                "Lead Score": lead.get("客户匹配评分", ""),
                "Company Domain": normalize_host(lead.get("公司官网", "")),
                "Country": lead.get("国家", ""),
                "Industry": lead.get("所属行业", ""),
            }
            for header, value in mappings.items():
                col = headers.index(header) + 1
                if not ws.cell(row, col).value:
                    ws.cell(row, col, value)
            discovery_angle = str(lead.get("为什么判断它可能需要 PCB/PCBA", "") or "").strip()
            angle_col = headers.index("Why This Angle Fits") + 1
            existing_angle = str(ws.cell(row, angle_col).value or "").strip()
            if "hardware signals:" in discovery_angle.lower() and "hardware signals:" not in existing_angle.lower():
                ws.cell(row, angle_col, " ".join(part for part in [existing_angle, discovery_angle] if part))

    def existing_company_keys(self) -> set[str]:
        ws = self.wb[LEADS_SHEET]
        headers = self.headers(ws)
        name_col = headers.index("公司名称") + 1
        website_col = headers.index("公司官网") + 1
        return {
            dedupe_key(ws.cell(row, name_col).value, ws.cell(row, website_col).value)
            for row in range(2, ws.max_row + 1)
        }

    def existing_domains(self) -> set[str]:
        ws = self.wb[LEADS_SHEET]
        headers = self.headers(ws)
        website_col = headers.index("公司官网") + 1
        return {
            domain
            for row in range(2, ws.max_row + 1)
            if (domain := normalize_domain(ws.cell(row, website_col).value))
        }

    def existing_emails(self) -> set[str]:
        emails = set()
        for sheet_name in [LEADS_SHEET, OUTREACH_SHEET, SENT_LOG_SHEET]:
            ws = self.wb[sheet_name]
            headers = self.headers(ws)
            for candidate in ["公开邮箱", "Contact Email"]:
                if candidate not in headers:
                    continue
                col = headers.index(candidate) + 1
                for row in range(2, ws.max_row + 1):
                    value = str(ws.cell(row, col).value or "").strip().lower()
                    if value and value not in {"not found", "未公开/待核实"}:
                        emails.add(value)
        return emails

    def sent_emails(self) -> set[str]:
        ws = self.wb[SENT_LOG_SHEET]
        headers = self.headers(ws)
        if "Contact Email" not in headers or "Status" not in headers:
            return set()
        email_col = headers.index("Contact Email") + 1
        status_col = headers.index("Status") + 1
        return {
            str(ws.cell(row, email_col).value or "").strip().lower()
            for row in range(2, ws.max_row + 1)
            if str(ws.cell(row, status_col).value or "").upper() == "SENT"
        }

    def sent_companies(self) -> set[str]:
        ws = self.wb[SENT_LOG_SHEET]
        headers = self.headers(ws)
        if "Company Name" not in headers or "Status" not in headers:
            return set()
        company_col = headers.index("Company Name") + 1
        status_col = headers.index("Status") + 1
        return {
            str(ws.cell(row, company_col).value or "").strip().lower()
            for row in range(2, ws.max_row + 1)
            if str(ws.cell(row, status_col).value or "").upper() == "SENT"
        }

    def sent_domains(self) -> set[str]:
        ws = self.wb[SENT_LOG_SHEET]
        headers = self.headers(ws)
        if "Company Domain" not in headers or "Status" not in headers:
            return set()
        domain_col = headers.index("Company Domain") + 1
        status_col = headers.index("Status") + 1
        email_col = headers.index("Contact Email") + 1 if "Contact Email" in headers else None
        source_col = headers.index("Source URL") + 1 if "Source URL" in headers else None
        domains: set[str] = set()
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, status_col).value or "").upper() != "SENT":
                continue
            raw_domain = ws.cell(row, domain_col).value
            domain = self._valid_host(raw_domain)
            if domain and registrable_domain(domain) and domain not in FREE_EMAIL_DOMAINS:
                domains.add(domain)
                continue
            if str(raw_domain or "").strip() or email_col is None or source_col is None:
                continue
            email_domain = self._valid_email_domain(ws.cell(row, email_col).value)
            if not email_domain or email_domain in FREE_EMAIL_DOMAINS:
                continue
            source_hosts = self._source_hosts(ws.cell(row, source_col).value)
            if any(self._hosts_are_attributable(email_domain, source_host) for source_host in source_hosts):
                domains.add(email_domain)
        return domains

    def append_leads(self, leads: list[Lead]) -> AppendStats:
        ws = self.wb[LEADS_SHEET]
        headers = self.headers(ws)
        keys = self.existing_company_keys()
        emails = self.existing_emails()
        added = skipped = score4 = score5 = 0
        for lead in leads:
            key = dedupe_key(lead.company_name, lead.website)
            email = lead.email.strip().lower()
            if key in keys or (email and email != "not found" and email in emails):
                skipped += 1
                continue
            keys.add(key)
            if email and email != "not found":
                emails.add(email)
            added += 1
            score4 += int(lead.score == 4)
            score5 += int(lead.score == 5)
            row_data = lead.to_excel_row(ws.max_row)
            ws.append([row_data.get(header, "") for header in headers])
        return AppendStats(added=added, skipped_duplicates=skipped, score4=score4, score5=score5)

    def append_outreach_drafts(self, drafts: list[dict[str, Any]]) -> int:
        ws = self.wb[OUTREACH_SHEET]
        headers = self.headers(ws)
        email_col = headers.index("Contact Email") + 1
        company_col = headers.index("Company Name") + 1
        domain_col = headers.index("Company Domain") + 1
        country_col = headers.index("Country") + 1
        existing_emails = {
            str(ws.cell(row, email_col).value or "").strip().lower()
            for row in range(2, ws.max_row + 1)
            if str(ws.cell(row, email_col).value or "").strip()
        }
        existing_domains = {
            domain
            for row in range(2, ws.max_row + 1)
            if (domain := normalize_host(ws.cell(row, domain_col).value))
        }
        existing_company_countries = {
            company_country_key(company, country)
            for row in range(2, ws.max_row + 1)
            if (company := str(ws.cell(row, company_col).value or "").strip())
            and (country := str(ws.cell(row, country_col).value or "").strip())
            and not is_generic_company_name(company)
        }
        sent_emails = self.sent_emails()
        sent_companies = self.sent_companies()
        sent_domains = self.sent_domains()
        added = 0
        for draft in drafts:
            company = str(draft.get("Company Name", "")).strip()
            email = str(draft.get("Contact Email", "")).strip().lower()
            domain = normalize_host(draft.get("Company Domain") or draft.get("Source URL"))
            country = str(draft.get("Country", "")).strip()
            generic_company = is_generic_company_name(company)
            company_country = company_country_key(company, country)
            if (
                (email and (email in existing_emails or email in sent_emails))
                or (domain and (domain in existing_domains or domain in sent_domains))
                or (
                    not generic_company
                    and company
                    and country
                    and (company_country in existing_company_countries or company.lower() in sent_companies)
                )
            ):
                continue
            draft.setdefault("Send Status", "")
            draft.setdefault("Send Date", "")
            draft.setdefault("Error Message", "")
            draft.setdefault("Do Not Contact", "")
            ws.append([draft.get(header, "") for header in headers])
            if email:
                existing_emails.add(email)
            if domain:
                existing_domains.add(domain)
            if not generic_company and company and country:
                existing_company_countries.add(company_country)
            added += 1
        self.evaluate_auto_send_all()
        return added

    def evaluate_auto_send_all(self) -> None:
        ws = self.wb[OUTREACH_SHEET]
        headers = self.headers(ws)
        sent_emails = self.sent_emails()
        sent_companies = self.sent_companies()
        for row in range(2, ws.max_row + 1):
            eligible, reason = self.evaluate_row(ws, headers, row, sent_emails, sent_companies)
            ws.cell(row, headers.index("Auto Send Eligible") + 1, "YES" if eligible else "NO")
            ws.cell(row, headers.index("Auto Send Reason") + 1, reason)

    @classmethod
    def final_send_safety_check(cls, ws: Worksheet, headers: list[str] | dict[str, int], row: int) -> tuple[bool, str]:
        """Require an attributable first-party company source before any outbound email."""
        def value(header: str) -> str:
            column = headers[header] if isinstance(headers, dict) else headers.index(header) + 1
            return str(ws.cell(row, column).value or "").strip()

        company = value("Company Name")
        company_host = cls._valid_host(value("Company Domain"))
        email_host = cls._valid_email_domain(value("Contact Email"))
        source = value("Source URL")
        source_urls = [url.rstrip(".,:)]}") for url in re.findall(r"https?://[^\s;]+", source, flags=re.IGNORECASE)]
        source_hosts = [cls._valid_host(url) for url in source_urls]
        source_text = " ".join([source, *source_urls]).casefold()
        company_name = re.sub(r"\s+", " ", company).strip().casefold()
        reasons: list[str] = []

        if not company_host and source_hosts:
            company_host = next((host for host in source_hosts if host), "")
        if not company_host:
            reasons.append("Company Website/domain missing or invalid")
        if not email_host or email_host in FREE_EMAIL_DOMAINS:
            reasons.append("Company email domain missing or free mailbox")
        elif not company_host or not cls._hosts_are_attributable(email_host, company_host):
            reasons.append("Email domain does not match company website")
        if not source_urls:
            reasons.append("Source URL missing or invalid")
        if company_host in _BLOCKED_SEND_HOSTS or company_host.startswith("wiki") or ".wiki" in company_host:
            reasons.append("Blocked encyclopedia/wiki company domain")
        if any(host in _BLOCKED_SEND_HOSTS or host.startswith("wiki") or ".wiki" in host for host in source_hosts):
            reasons.append("Blocked encyclopedia/wiki source")
        if any(term in source_text for term in _BLOCKED_SOURCE_TERMS):
            reasons.append("Blocked directory, list, public-sector, blog, media, or service source")
        if company_name in _BLOCKED_COMPANY_NAMES or is_generic_company_name(company):
            reasons.append("Company name is generic or page-title derived")
        if _PAGE_TITLE_COMPANY_NAME_RE.search(company):
            reasons.append("Company name appears to be a page title or sentence, not a legal company identity")
        # A directory/exhibitor page may appear in the trace, but it cannot be the
        # sole evidence. At least one URL must resolve to the company-controlled host.
        if source_hosts and company_host and not any(cls._hosts_are_attributable(host, company_host) for host in source_hosts):
            reasons.append("Source URL is not an attributable company official page")

        return (False, "; ".join(dict.fromkeys(reasons))) if reasons else (True, "Passed final company, source, and email-domain safety check")

    def evaluate_row(self, ws: Worksheet, headers: list[str], row: int, sent_emails: set[str], sent_companies: set[str]) -> tuple[bool, str]:
        def value(header: str) -> str:
            return str(ws.cell(row, headers.index(header) + 1).value or "").strip()

        reasons = []
        score = int(float(value("Lead Score") or 0))
        email = value("Contact Email")
        email_type = value("Email Type")
        company = value("Company Name")
        body = value("Email Body")
        subject = value("Subject")
        source = value("Source URL")
        compliance = value("Compliance Note")
        status = value("Send Status")
        dnc = value("Do Not Contact").upper()
        angle = value("Why This Angle Fits")
        word_count = len(body.replace("/", " ").split())
        if not email or email == "Not Found":
            reasons.append("Email missing")
        if email and email != "Not Found" and not normalize_mailbox(email):
            reasons.append("Invalid email format")
        if email_type not in _ALLOWED_SEND_EMAIL_TYPES:
            reasons.append(f"Email Type not allowed: {email_type or 'blank'}")
        email_domain = self._valid_email_domain(email)
        if email_domain in FREE_EMAIL_DOMAINS:
            reasons.append("Free email domain not allowed")
        elif email and email != "Not Found" and not self._recipient_is_attributable(
            email,
            value("Company Domain"),
            source,
        ):
            reasons.append("Recipient domain not attributable to company domain or official source URL")
        if not source:
            reasons.append("Source URL missing")
        if not compliance:
            reasons.append("Compliance Note missing")
        if dnc == "YES":
            reasons.append("Do Not Contact")
        if status:
            reasons.append(f"Send Status is {status}")
        if email.lower() in sent_emails:
            reasons.append("Email already sent")
        if company.lower() in sent_companies:
            reasons.append("Company already sent")
        if company.casefold() in _UNRESOLVED_COMPANY_NAMES or is_generic_company_name(company):
            reasons.append("Company name unresolved or generic")
        source_urls = re.findall(r"https?://[^\s;]+", source, flags=re.IGNORECASE)
        if source_urls and all(is_directory_url(url) for url in source_urls):
            reasons.append("Source is a directory operator")
        recorded_signals = self._recorded_hardware_signals(angle)
        qualification = hardware_qualification(list(recorded_signals))
        exclusion_signals = hardware_qualification(
            [angle, source],
            company_name=company,
        ).exclusion_signals
        if exclusion_signals:
            reasons.append(f"Excluded prospect type: {', '.join(exclusion_signals)}")
        if not qualification.eligible:
            reasons.append("Hardware qualification failed")
        if len(recorded_signals) < 2:
            reasons.append("Fewer than two recorded hardware signals")
        if OPT_OUT_SENTENCE not in body:
            reasons.append("Opt-out sentence missing")
        if not subject:
            reasons.append("Subject missing")
        if not body:
            reasons.append("Email Body missing")
        if word_count < 120 or word_count > 180:
            reasons.append(f"Word count out of range: {word_count}")
        if is_likely_peer_manufacturer(" ".join([company, value("Why This Angle Fits"), body])):
            reasons.append("Possible PCB/PCBA peer")
        final_safe, final_reason = self.final_send_safety_check(ws, headers, row)
        ws.cell(row, headers.index("Final Send Safety Check") + 1, "PASS" if final_safe else "BLOCKED")
        ws.cell(row, headers.index("Final Send Block Reason") + 1, "" if final_safe else final_reason)
        if not final_safe:
            reasons.append(final_reason)
        return (False, "; ".join(reasons)) if reasons else (True, "Eligible for automatic sending")

    @staticmethod
    def _recorded_hardware_signals(angle: str) -> tuple[str, ...]:
        recorded_evidence = _RECORDED_HARDWARE_SIGNALS_RE.findall(angle or "")
        return hardware_qualification(recorded_evidence).hardware_signals

    def pending_auto_eligible_count(self) -> int:
        ws = self.wb[OUTREACH_SHEET]
        headers = self.headers(ws)
        eligible_col = headers.index("Auto Send Eligible") + 1
        status_col = headers.index("Send Status") + 1
        return sum(
            1
            for row in range(2, ws.max_row + 1)
            if str(ws.cell(row, eligible_col).value or "").upper() == "YES"
            and not str(ws.cell(row, status_col).value or "").strip()
        )

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.path)

    def append_sent_log(self, company: str, email: str, subject: str, status: str, error: str = "", country: str = "", industry: str = "", source_url: str = "", domain: str = "") -> None:
        ws = self.wb[SENT_LOG_SHEET]
        ws.append([datetime.now().isoformat(timespec="seconds"), company, email, subject, status, error, country, industry, source_url, domain])
